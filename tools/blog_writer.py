from __future__ import annotations
import io
import re
import streamlit as st
from docx import Document
from utils import history as history_utils
from utils import pipeline, feedback, rag, metrics
from utils import ui as ui_utils


def _add_paragraph_with_bold(doc: Document, text: str, style: str | None = None) -> None:
    """**囲みを太字runとして段落に追加する"""
    para = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    parts = re.split(r"\*\*", text)
    for i, part in enumerate(parts):
        run = para.add_run(part)
        run.bold = (i % 2 == 1)


def _to_docx(markdown_text: str) -> bytes:
    doc = Document()
    for line in markdown_text.splitlines():
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("- ") or line.startswith("* "):
            _add_paragraph_with_bold(doc, line[2:].strip(), style="List Bullet")
        elif re.match(r"^\d+\. ", line):
            _add_paragraph_with_bold(doc, re.sub(r"^\d+\. ", "", line).strip(), style="List Number")
        elif line.strip():
            _add_paragraph_with_bold(doc, line.strip())
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


TONES = {
    "プロフェッショナル": "堅実で信頼感のある、情報重視の文体",
    "フレンドリー": "フレンドリーでカジュアル、親しみやすい文体",
    "エンタメ系": "楽しく引き込む、軽快な文体",
    "教育的": "教育的で分かりやすい、順を追った文体",
}

LENGTHS = {
    "短め (500字程度)": 500,
    "標準 (1000字程度)": 1000,
    "詳細 (2000字程度)": 2000,
    "しっかり (2500字程度)": 2500,
    "大ボリューム (3000字程度)": 3000,
}


def _score_dims(report) -> list[tuple[str, int]]:
    """LLM 評価スコアを内訳バー用の (ラベル, 値) に変換する。"""
    llm = report.details.get("llm") or {}
    if not llm:
        return []
    return [
        ("AI感の低さ", int(llm.get("ai_feel", 0))),
        ("読者への共感", int(llm.get("empathy", 0))),
        ("具体性・体験談", int(llm.get("concrete", 0))),
    ]


def _guardrail_panel(report) -> None:
    ui_utils.render_score_panel(report.score, report.passed, _score_dims(report))
    if report.violations or report.warnings:
        with st.expander("判定の詳細", expanded=not report.passed):
            for v in report.violations:
                st.markdown(f"- ❌ {v}")
            for w in report.warnings:
                st.markdown(f"- ⚠️ {w}")


def _feedback_section(saved: dict) -> None:
    """Task 3: ユーザー修正差分・採否フィードバックの蓄積"""
    st.divider()
    st.subheader("フィードバック（学習用ログ）")
    st.caption(
        "本文を手直しして「採用」または「ボツ」を押すと、編集差分が logs/feedback.jsonl "
        "に保存され、次回以降のお手本（Few-Shot）として使われます。"
    )

    original = saved["final"]
    edited = st.text_area(
        "本文（自由に手直しできます）",
        value=st.session_state.get("blog_writer_edited", original),
        height=320,
        key="blog_writer_edit_area",
    )
    st.session_state["blog_writer_edited"] = edited

    if saved.get("feedback_saved"):
        st.info("この記事のフィードバックは保存済みです。", icon=":material/save:")

    col_a, col_b = st.columns(2)
    with col_a:
        adopt = st.button("採用（この方向で良い）", icon=":material/thumb_up:", use_container_width=True, type="primary")
    with col_b:
        reject = st.button("ボツ（作り直したい）", icon=":material/thumb_down:", use_container_width=True)

    if adopt or reject:
        score = feedback.ADOPTED if adopt else feedback.REJECTED
        feedback.save_feedback(
            original_output=original,
            user_edited_output=edited,
            feedback_score=score,
            keywords=saved.get("keywords", ""),
            meta={"tool": "blog_writer", "title": saved.get("title", ""), "retries": saved.get("retries", 0)},
        )
        saved["feedback_saved"] = True
        st.session_state["blog_writer_result"] = saved
        st.success("フィードバックを保存しました。ありがとうございます！")

    s = feedback.stats()
    st.caption(f"蓄積ログ: 全{s['total']}件（採用 {s['adopted']} / ボツ {s['rejected']} / 編集あり {s['edited']}）")


def render():
    st.header(":material/edit: ブログ記事作成")
    st.caption("Planner → Writer → Polisher の3段階で、AI感を抑えた記事を生成します")

    col1, col2 = st.columns(2)
    with col1:
        title = st.text_input("記事タイトル / テーマ *", placeholder="例: 初心者向けPython入門ガイド")
        tone = st.selectbox("文体・トーン", list(TONES.keys()), index=1)
    with col2:
        keywords = st.text_input("キーワード（カンマ区切り）", placeholder="例: Python, プログラミング, 入門")
        length = st.selectbox("文章量", list(LENGTHS.keys()))

    target = st.text_input("ターゲット読者", placeholder="例: プログラミング未経験の社会人")
    extra = st.text_area("追加の指示・補足（任意）", placeholder="例: SEOを意識した構成にしてください", height=80, max_chars=300)

    col_opt1, col_opt2 = st.columns(2)
    with col_opt1:
        use_fewshot = st.checkbox("過去の採用例をお手本にする（Few-Shot RAG）", value=True)
    with col_opt2:
        max_retries = st.slider("ガードレール未達時の再校正回数", 0, 3, 1)

    if st.button("記事を生成", type="primary", use_container_width=True):
        if not title:
            st.warning("記事タイトル / テーマを入力してください")
            return
        if len(title) > 200:
            st.warning("タイトルが長すぎます（200字以内）")
            return

        fewshot = rag.fewshot_for(theme=title, keywords=keywords) if use_fewshot else ""

        status = st.status("生成パイプラインを実行中...", expanded=True)

        def _progress(msg: str) -> None:
            status.write(msg)

        try:
            result = pipeline.run_pipeline(
                theme=title,
                target=target,
                keywords=keywords,
                char_count=LENGTHS[length],
                tone_ja=TONES[tone],
                extra=extra,
                fewshot=fewshot,
                max_retries=max_retries,
                progress=_progress,
            )
        except RuntimeError as e:
            status.update(label="生成に失敗しました", state="error")
            st.error(str(e))
            return

        status.update(label="生成完了", state="complete")

        if fewshot:
            status.write("お手本として過去の採用例を参照しました")

        report = result["guardrail"]
        metrics.log_generation(
            tool="ブログ記事作成",
            title=title,
            score=report.score,
            passed=report.passed,
            retries=result["retries"],
            dims=(report.details.get("llm") or {}),
        )

        st.session_state["blog_writer_result"] = {
            "content": result["final"],
            "final": result["final"],
            "draft": result["draft"],
            "plan": result["plan"],
            "guardrail": result["guardrail"],
            "retries": result["retries"],
            "title": title,
            "keywords": keywords,
            "used_fewshot": bool(fewshot),
            "feedback_saved": False,
        }
        st.session_state["blog_writer_edited"] = result["final"]

    # 生成結果の表示
    saved = st.session_state.get("blog_writer_result")
    if saved:
        st.divider()
        st.subheader("生成結果")

        _guardrail_panel(saved["guardrail"])
        if saved.get("retries"):
            st.caption(f"ガードレール未達のため {saved['retries']} 回、自動で再校正しました。")
        if saved.get("used_fewshot"):
            st.caption("過去の採用例をお手本に反映しています。")

        with st.expander("① 構成・共感設計（Planner の出力）"):
            st.markdown(saved["plan"])
        with st.expander("② 初稿（Writer の出力・校正前）"):
            st.markdown(saved["draft"])

        st.markdown("#### ③ 完成稿（Polisher の出力）")
        st.markdown(saved["final"])

        export_text = st.session_state.get("blog_writer_edited", saved["final"])
        col_dl1, col_dl2, col_save = st.columns(3)
        with col_dl1:
            st.download_button(
                "テキスト (.txt)",
                data=export_text,
                file_name="blog_article.txt",
                mime="text/plain",
                icon=":material/description:",
                use_container_width=True,
            )
        with col_dl2:
            st.download_button(
                "Word (.docx)",
                data=_to_docx(export_text),
                file_name="blog_article.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                icon=":material/article:",
                use_container_width=True,
            )
        with col_save:
            if st.button("保存", icon=":material/save:", use_container_width=True):
                history_utils.save_entry("ブログ記事作成", saved["title"], export_text)
                st.success("保存しました！")

        _feedback_section(saved)
